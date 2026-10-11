import openpyxl
from docx import Document

def save_to_xlsx(data: dict, filename: str = "report.xlsx"):
    """Сохраняет словарь с результатами в Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Результаты расчета"
    
    ws.append(["Параметр", "Значение"])
    for key, value in data.items():
        ws.append([key, value])
        
    wb.save(filename)
    print(f"Отчет успешно сохранен в {filename}")

def save_to_docx(data: dict, filename: str = "report.docx"):
    """Сохраняет словарь с результатами в Word."""
    doc = Document()
    doc.add_heading('Отчет по расчету геометрического тела', 0)
    
    for key, value in data.items():
        doc.add_paragraph(f"{key}: {value}")
        
    doc.save(filename)
    print(f"Отчет успешно сохранен в {filename}")